# gen_docs_kiia.py — 부록 E 실습 패키지: KIIA-1·3·5·8 DOCX 문서 생성
# 전량 가상 데이터. 책(ch06~09·11·13) 내용과 일치. 첫 줄 고지 문구 삽입.
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = r"C:\Users\user\Documents\developer\7_audit\practice-package"
NOTICE = "본 자료는 실습용 가상 데이터이며 실재 인물·기관·거래와 무관하다."
FONT = "맑은 고딕"
NAVY = RGBColor(0x1A, 0x2E, 0x5A)
GREY = RGBColor(0x55, 0x55, 0x55)

# ── 공통 헬퍼 ────────────────────────────────────────────────
def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for sec in doc.sections:
        sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.2)
        sec.left_margin = Cm(2.4); sec.right_margin = Cm(2.4)
    return doc

def _set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color is not None: run.font.color.rgb = color
    if italic is not None: run.font.italic = italic

def title(doc, text, sub=None):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); _set_font(r, 17, True, NAVY)
    p.paragraph_format.space_after = Pt(4)
    if sub:
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(sub); _set_font(r2, 11, False, GREY)
        p2.paragraph_format.space_after = Pt(8)

def h1(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); _set_font(r, 13, True, NAVY)
    return p

def h2(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); _set_font(r, 11.5, True)
    return p

def para(doc, text, size=10.5, bold=False, color=None, align=None, space=4, italic=False):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space)
    r = p.add_run(text); _set_font(r, size, bold, color, italic)
    return p

def notice(doc):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(10)
    r = p.add_run(NOTICE); _set_font(r, 9.5, False, GREY, italic=True)

def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); _set_font(r, size)
    return p

def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def table(doc, headers, rows, widths=None, header_fill="1A2E5A", body_size=9.5, notice_row=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        _shade(hdr[i], header_fill)
        pp = hdr[i].paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.paragraph_format.space_after = Pt(1)
        r = pp.add_run(str(htext)); _set_font(r, body_size+0.5, True, RGBColor(0xFF,0xFF,0xFF))
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            pp = cells[i].paragraphs[0]; pp.paragraph_format.space_after = Pt(1)
            r = pp.add_run("" if val is None else str(val)); _set_font(r, body_size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def save(doc, folder, fname):
    path = os.path.join(ROOT, folder, fname)
    doc.save(path)
    print("생성:", os.path.join(folder, fname))

# ═══════════════════════════════════════════════════════════
# KIIA-1  완성 감사보고서
# ═══════════════════════════════════════════════════════════
def kiia1():
    doc = new_doc()
    title(doc, "2026년 1분기 구매·경비 정기감사 결과보고서", "ABC 제조(주) 감사실")
    notice(doc)
    para(doc, "본 보고서는 실습용 정답지로 작성한 완성형 감사결과보고서다. 발견사항은 기준·현황·원인·결과·개선권고의 5C 구조로 기술하며, 보고 품질은 국제내부감사기준(GIAS) 표준 11.2의 일곱 특성을 준거로 한다.", size=10, color=GREY)

    h1(doc, "1. 감사 개요")
    table(doc, ["구분", "내용"], [
        ["감사명", "2026년 1분기 구매·경비 정기감사"],
        ["감사 근거", "ABC 제조(주) 내부감사규정 제8조(정기감사)"],
        ["감사 대상", "구매팀·총무팀·영업팀의 2026년 1분기(1~3월) 구매·경비 집행"],
        ["감사 기간", "2026.04.06 ~ 2026.04.24"],
        ["감사반", "감사팀장 김감사 외 감사역 2명"],
        ["보고일", "2026.05.08"],
    ], widths=[3.2, 12.5])

    h1(doc, "2. 감사 범위 및 방법")
    para(doc, "감사 대상은 위험 기반으로 선정하였다. 전분기 지적 이력, 거래 규모, 통제 취약성을 기준으로 구매·경비 영역을 우선순위에 두었다. 분석은 전수 점검과 표본 점검을 병행하였다.")
    table(doc, ["방법", "적용 대상", "내용"], [
        ["전수 점검", "법인카드 사용내역·수의계약 대장", "1분기 전건에 이상거래 규칙을 일괄 적용"],
        ["벤포드 분석", "경비 지출 금액", "첫째 자리 숫자 분포를 기대치와 대조해 확인 대상 선별"],
        ["표본 점검", "접대비 정산 증빙", "위험 기준 표본을 추출해 증빙 요건 대조"],
    ], widths=[3.0, 5.2, 7.6])
    para(doc, "벤포드 분석과 규칙 기반 선별 결과는 확인 대상을 좁히는 신호이며, 부적정 여부는 감사인이 원자료·증빙과 대조해 확정하였다.", size=10, color=GREY)

    h1(doc, "3. 발견사항")
    # 발견 1
    h2(doc, "발견 1. 유사 용역의 분할 수의계약으로 인한 경쟁입찰 회피")
    table(doc, ["5C", "내용"], [
        ["기준(Criteria)", "구매규정 제15조: 동일·유사 용역의 합산 금액이 5,000만원 이상이면 경쟁입찰을 원칙으로 한다. 제16조: 목적·대상이 유사한 반복 용역은 합산 검토한다."],
        ["현황(Condition)", "구매팀은 물류 위탁 용역 3건을 동일 업체(세종물류(주))와 각 4,700만·4,800만·4,900만원에 수의계약으로 체결하였다. 계약은 1~3월 약 1개월 간격이며 합산 시 1억 4,400만원이다. 경쟁입찰 검토 기록은 확인되지 않았다."],
        ["원인(Cause)", "유사 용역의 분기 합산 점검이 계약 건별 처리 절차에 반영되지 않았고, 반복 발주를 합산해 경보하는 통제가 없었다."],
        ["결과(Consequence)", "경쟁입찰을 통한 가격 검증 기회가 상실되고 특정 업체 편중이 발생하였다."],
        ["개선권고(Corrective action)", "동일·유사 용역의 분기 합산 점검 절차를 신설하고, 수의계약 대장에 합산 한도 경보 기준을 반영한다."],
    ], widths=[3.6, 12.1])
    # 발견 2
    h2(doc, "발견 2. 법인카드 제한 업종·비업무시간 사용")
    table(doc, ["5C", "내용"], [
        ["기준(Criteria)", "법인카드 관리규정 제7조: 유흥업종 등 제한 업종 사용을 금지한다. 제9조: 사용 목적을 사전 등록하고 업무시간 내 사용을 원칙으로 한다."],
        ["현황(Condition)", "1분기 법인카드 사용내역 전수 점검 결과, 유흥업종 결제 4건과 심야·주말 사용 9건이 확인되었다. 해당 거래에는 사용 목적 사전 등록과 증빙이 미비하였다."],
        ["원인(Cause)", "제한 업종 결제를 차단·경보하는 통제가 없었고, 사용 목적의 사후 점검이 형식에 그쳤다."],
        ["결과(Consequence)", "업무 무관 지출 위험이 상존하고 경비 통제의 신뢰성이 저하되었다."],
        ["개선권고(Corrective action)", "제한 업종 결제 자동 차단·경보를 도입하고, 심야·주말 사용에 대한 사유 등록을 의무화한다."],
    ], widths=[3.6, 12.1])
    # 발견 3
    h2(doc, "발견 3. 접대비 증빙 미비")
    table(doc, ["5C", "내용"], [
        ["기준(Criteria)", "경비처리규정 제12조: 접대비는 참석자·목적을 기재하고 적격증빙을 첨부한다."],
        ["현황(Condition)", "접대비 표본 20건 중 6건에서 참석자·목적 기재가 누락되었고, 이 중 3건은 적격증빙이 첨부되지 않았다."],
        ["원인(Cause)", "증빙 요건 안내와 정산 단계 검증이 미흡하였다."],
        ["결과(Consequence)", "손금 부인 위험이 있고 지출의 업무 관련성을 입증하기 어렵다."],
        ["개선권고(Corrective action)", "접대비 정산 시 참석자·목적을 필수 입력 항목으로 두고, 적격증빙 확인 절차를 정산 통제에 포함한다."],
    ], widths=[3.6, 12.1])

    h1(doc, "4. 리스크 요약")
    para(doc, "발견사항의 위험등급은 영향도와 발생 가능성을 각각 세 단계로 평가해 부여하였다.")
    table(doc, ["발견", "영향도", "발생 가능성", "위험등급"], [
        ["발견 1 분할 수의계약", "높음", "중간", "High"],
        ["발견 2 법인카드 부적절 사용", "중간", "높음", "High"],
        ["발견 3 접대비 증빙 미비", "중간", "중간", "Medium"],
    ], widths=[6.6, 3.0, 3.2, 2.9])

    h1(doc, "5. 개선 권고")
    table(doc, ["권고", "담당 부서", "이행 기한"], [
        ["유사 용역 분기 합산 점검 절차 신설 및 대장 경보 반영", "구매팀", "2026.06.30"],
        ["제한 업종 자동 차단·경보 및 심야·주말 사용 사유 등록 의무화", "총무팀", "2026.07.31"],
        ["접대비 정산 필수 입력 항목 및 적격증빙 확인 통제 보완", "영업팀·재무팀", "2026.06.30"],
    ], widths=[9.0, 3.5, 3.2])

    h1(doc, "6. 피감부서 의견")
    para(doc, "구매팀은 합산 점검 절차의 필요성에 동의하며 대장 시스템 반영을 6월 중 완료하겠다는 의견을 제시하였다. 총무팀은 제한 업종 차단 기능의 시스템 개발 일정을 고려해 이행 기한 연장을 협의 요청하였다. 영업팀은 증빙 요건을 재안내하고 정산 통제 보완에 협조하겠다고 회신하였다.")

    h1(doc, "7. 결론")
    para(doc, "1분기 구매·경비 집행에서 분할 수의계약, 법인카드 부적절 사용, 접대비 증빙 미비 세 건이 확인되었다. 세 건 모두 통제 절차의 공백 또는 실효성 부족에서 비롯되었으며, 개선권고는 확인된 원인에 대응하도록 작성하였다. 본 보고서의 개선권고는 사후관리 대장으로 이관하여 이행 기한, 담당 부서, 증빙, 상태를 기준으로 상시 점검한다. 조치 완료 회신과 개선 효과 확인은 구분하여, 통제 신설 이후의 실제 기록으로 효과를 검증한다.")

    save(doc, "KIIA-1_완성_감사보고서", "sample_audit_report.docx")

# ═══════════════════════════════════════════════════════════
# KIIA-3  2026년 2분기 감사계획서 (대외비)
# ═══════════════════════════════════════════════════════════
def kiia3():
    doc = new_doc()
    para(doc, "대외비", size=10.5, bold=True, color=RGBColor(0xB0,0x00,0x00), align=WD_ALIGN_PARAGRAPH.RIGHT, space=2)
    title(doc, "2026년 2분기 감사계획서", "ABC 제조(주) 감사실")
    notice(doc)

    h1(doc, "1. 감사 목적")
    para(doc, "2026년 2분기 정기감사는 전분기 지적사항의 이행 여부를 확인하고, 위험 기반으로 선정한 부서의 업무 처리와 내부통제 준수 여부를 점검하는 데 목적이 있다. 감사 결과는 개선권고와 사후관리 대장으로 연결한다.")

    h1(doc, "2. 감사 대상 및 우선순위")
    para(doc, "위험 기반으로 다섯 개 부서를 감사 대상으로 선정하였다. 우선순위는 전분기 지적 이력, 거래 규모, 통제 취약성을 기준으로 부여한다. 일부 부서는 우선순위 근거가 기재되지 않았으므로, 실습에서 위험평가 기준에 따라 근거를 보완한다.")
    table(doc, ["우선순위", "부서", "주요 점검 영역", "선정 근거"], [
        ["1", "구매팀", "수의계약·분할발주·협력업체 등록", "전분기 분할 수의계약 지적의 후속 점검"],
        ["2", "IT운영팀", "접근권한·외주용역·장애대응", "접근권한 관리 및 외주 계약 통제 위험"],
        ["3", "총무팀", "법인카드·비품·차량 관리", ""],
        ["4", "인사팀", "퇴직자 계정·보안교육·급여", "퇴직자 계정 회수 지연 위험"],
        ["5", "영업지원팀", "접대비·출장비·매출채권", ""],
    ], widths=[2.0, 2.6, 6.4, 4.7])
    para(doc, "※ 3순위 총무팀과 5순위 영업지원팀은 선정 근거가 비어 있다. 위험평가 기준을 적용해 근거를 작성하는 것이 실습 과제다.", size=9.5, color=GREY)

    h1(doc, "3. 감사 방법론")
    table(doc, ["방법론", "적용 내용"], [
        ["위험 기반 감사", "위험평가 결과에 따라 감사 대상과 표본을 우선순위화한다."],
        ["벤포드 분석", "경비·구매 금액의 첫째 자리 분포를 기대치와 대조해 확인 대상을 선별한다."],
        ["연속감사(continuous auditing)", "법인카드·경비 데이터에 이상거래 규칙을 상시 적용해 예외를 조기에 감지한다."],
    ], widths=[4.6, 11.1])

    h1(doc, "4. 감사 일정")
    table(doc, ["단계", "기간", "주요 활동"], [
        ["사전 준비", "2026.06.15 ~ 06.19", "위험평가, 자료 요청, 감사 프로그램 확정"],
        ["현장 감사", "2026.06.22 ~ 07.03", "부서별 자료 점검, 인터뷰, 이상거래 분석"],
        ["정리·검토", "2026.07.06 ~ 07.10", "발견사항 정리, 조서 검토, 피감의견 청취"],
        ["보고", "2026.07.13 ~ 07.17", "감사결과보고서 작성 및 결재"],
    ], widths=[3.0, 4.6, 8.1])

    h1(doc, "5. 투입 자원")
    table(doc, ["구분", "내용"], [
        ["감사 인력", "감사팀장 1명, 감사역 3명"],
        ["분석 도구", "생성형 AI 기반 문서 검토·데이터 분석(조직 승인 환경)"],
        ["예상 공수", "총 40인일"],
    ], widths=[3.4, 12.3])

    h1(doc, "6. 예산")
    table(doc, ["항목", "금액(원)", "비고"], [
        ["감사 활동비", "3,000,000", "현장 감사 및 자료 확보"],
        ["분석 도구 이용료", "1,200,000", "AI 도구 유료 등급(분기)"],
        ["예비비", "800,000", "-"],
        ["합계", "5,000,000", ""],
    ], widths=[5.0, 4.0, 6.7])

    save(doc, "KIIA-3_감사계획서", "audit_plan_2026_Q2.docx")

# ═══════════════════════════════════════════════════════════
# KIIA-5  부서별 세부 감사계획서 5종 + 표준서식
# ═══════════════════════════════════════════════════════════
DEPT_PLANS = {
    "dept_audit_general.docx": {
        "dept": "총무팀",
        "bg": "총무팀은 법인카드, 비품, 차량, 일반 경비를 관리한다. 전분기 법인카드 부적절 사용이 지적되어 통제 실효성 점검이 필요하다.",
        "scope": "2026년 1~6월 법인카드 사용내역, 비품 구매·불용 처리, 차량 운행·유지 경비",
        "proc": [
            ["법인카드 전수 점검", "이상거래 규칙(심야·주말·제한업종·고액·분할) 적용", "확인 대상 거래 목록"],
            ["비품 관리 대조", "구매·불용 대장과 실물 대조 표본 점검", "불일치 목록"],
            ["차량 경비 분석", "유류·정비 지출의 벤포드·이상치 분석", "이상 지출 후보"],
        ],
        "risk": [
            ["법인카드 제한 업종·비업무시간 사용", "높음", "높음", "High"],
            ["비품 대장과 실물 불일치", "중간", "중간", "Medium"],
            ["차량 경비 과다·중복 청구", "중간", "중간", "Medium"],
        ],
    },
    "dept_audit_hr.docx": {
        "dept": "인사팀",
        "bg": "인사팀은 채용, 급여, 계정 권한, 보안교육을 담당한다. 퇴직자 계정 회수 지연과 보안교육 미이행이 반복 위험으로 확인된다.",
        "scope": "2026년 1~6월 급여·초과근무 지급, 퇴직자 계정 회수, 신규 입사자 보안교육 이행",
        "proc": [
            ["급여 데이터 분석", "직급·부서별 기본급·초과근무수당 이상치(Z점수·IQR) 및 동일 계좌 중복 탐지", "이상치·중복 계좌 후보"],
            ["퇴직자 계정 점검", "퇴직일과 계정 회수일 대조", "회수 지연 목록"],
            ["보안교육 이행 점검", "신규 입사자 명단과 교육 이수 기록 대조", "미이행 목록"],
        ],
        "risk": [
            ["퇴직자 계정 회수 지연", "높음", "높음", "High"],
            ["동일 계좌 중복 급여 수령(유령직원 의심)", "높음", "낮음", "Medium"],
            ["신규 입사자 보안교육 미이행", "중간", "중간", "Medium"],
        ],
    },
    "dept_audit_it.docx": {
        "dept": "IT운영팀",
        "bg": "IT운영팀은 시스템 접근권한, 장애 대응, 외주 용역, 정보자산을 관리한다. 접근권한 과다 부여와 장애 보고 지연이 주요 위험이다.",
        "scope": "2026년 1~6월 시스템 접근권한 부여·회수, 장애 대응·보고, 외주 용역 계약",
        "proc": [
            ["접근권한 점검", "권한 부여 대장과 재직·직무 대조, 과다 권한 식별", "과다 권한 목록"],
            ["장애 대응 점검", "장애 인지·보고·등록 시점과 규정 절차(To-Be) 대비", "지연·미등록 목록"],
            ["외주 계약 점검", "외주 용역의 분할·수의계약 여부, 산출물 검수 기록 확인", "계약 확인 대상"],
        ],
        "risk": [
            ["시스템 접근권한 과다 부여·회수 지연", "높음", "중간", "High"],
            ["장애 보고·등록 지연", "중간", "높음", "High"],
            ["외주 용역 검수 절차 미흡", "중간", "중간", "Medium"],
        ],
    },
    "dept_audit_purchase.docx": {
        "dept": "구매팀",
        "bg": "구매팀은 계약, 수의계약, 협력업체 등록, 검수를 담당한다. 전분기 유사 용역 분할 수의계약이 지적되어 후속 점검이 필요하다.",
        "scope": "2026년 1~6월 수의계약·경쟁입찰, 협력업체 등록·평가, 납품 검수",
        "proc": [
            ["수의계약 합산 점검", "동일·유사 용역의 분기 합산과 경쟁입찰 회피 여부 확인", "분할 의심 계약 목록"],
            ["협력업체 등록 점검", "신규 협력업체 사전 실사·등록 절차 이행 확인", "미실사 등록 목록"],
            ["검수 기록 점검", "납품 검수 기록과 대금 지급의 정합성 대조", "검수 누락 목록"],
        ],
        "risk": [
            ["유사 용역 분할 수의계약", "높음", "높음", "High"],
            ["협력업체 사전 실사 미흡", "중간", "중간", "Medium"],
            ["검수 없는 대금 지급", "높음", "낮음", "Medium"],
        ],
    },
    "dept_audit_sales.docx": {
        "dept": "영업지원팀",
        "bg": "영업지원팀은 접대비, 출장비, 매출채권을 관리한다. 접대비 증빙 미비와 매출채권 관리 지연이 주요 위험이다.",
        "scope": "2026년 1~6월 접대비 정산, 출장비 지급, 매출채권 회수",
        "proc": [
            ["접대비 증빙 점검", "참석자·목적 기재와 적격증빙 첨부 여부 표본 점검", "증빙 미비 목록"],
            ["출장비 분석", "출장 신청·정산 대조 및 중복·과다 지급 이상치 분석", "이상 지급 후보"],
            ["매출채권 점검", "채권 연령 분석과 회수 지연·대손 처리 적정성 확인", "장기 미회수 목록"],
        ],
        "risk": [
            ["접대비 증빙 미비", "중간", "높음", "High"],
            ["출장비 중복·과다 지급", "중간", "중간", "Medium"],
            ["매출채권 회수 지연", "중간", "중간", "Medium"],
        ],
    },
}

def kiia5_plan(fname, spec):
    doc = new_doc()
    title(doc, f"{spec['dept']} 세부 감사계획서", "ABC 제조(주) 감사실 · 2026년 2분기")
    notice(doc)

    h1(doc, "1. 배경")
    para(doc, spec["bg"])

    h1(doc, "2. 감사 범위")
    para(doc, spec["scope"])

    h1(doc, "3. 감사 절차")
    table(doc, ["절차", "방법", "산출물"], spec["proc"], widths=[4.2, 7.5, 4.0])

    h1(doc, "4. 리스크 매트릭스")
    table(doc, ["위험 요인", "영향도", "발생 가능성", "위험등급"], spec["risk"], widths=[8.0, 2.6, 2.8, 2.3])
    para(doc, "위험등급은 영향도와 발생 가능성을 각각 세 단계로 평가해 High·Medium·Low로 부여한다.", size=9.5, color=GREY)

    h1(doc, "5. 일정")
    table(doc, ["단계", "기간", "주요 활동"], [
        ["자료 확보", "2026.06.15 ~ 06.19", "감사 자료 요청 및 위험평가"],
        ["현장 점검", "2026.06.22 ~ 06.30", f"{spec['dept']} 자료 점검 및 인터뷰"],
        ["정리·보고", "2026.07.01 ~ 07.10", "발견사항 정리 및 부서 계획 보고"],
    ], widths=[3.0, 4.6, 8.1])

    save(doc, "KIIA-5_5개부서_감사", fname)

def kiia5_form():
    doc = new_doc()
    title(doc, "감사결과보고서", "표준서식(빈 양식)")
    notice(doc)
    para(doc, "본 서식은 공공감사에 관한 법률 시행규칙의 취지를 반영한 감사결과보고서 빈 표준서식이다. 각 빈칸을 실습에서 감사 사실과 처분요구 문안으로 채운다.", size=10, color=GREY)

    def blank(label, hint=""):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        r = p.add_run(f"{label}: "); _set_font(r, 10.5, True)
        r2 = p.add_run("[                                                        ]"); _set_font(r2, 10.5)
        if hint:
            r3 = p.add_run(f"  {hint}"); _set_font(r3, 9, False, GREY)

    h1(doc, "1. 감사 개요")
    blank("기관명"); blank("감사부서"); blank("담당자")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run("감사기간: "); _set_font(r, 10.5, True)
    r2 = p.add_run("[            ]  ~  [            ]"); _set_font(r2, 10.5)

    h1(doc, "2. 감사 목적")
    para(doc, "[                                                                                    ]", space=10)

    h1(doc, "3. 관계 법령 및 내부 기준")
    para(doc, "[                                                                                    ]", space=4)
    para(doc, "[                                                                                    ]", space=10)

    h1(doc, "4. 감사결과(확인된 사항)")
    para(doc, "관계 법령·내부 기준이 기대한 행위와 실제 확인된 행위를 차례로 제시하고, 기준·사실·차이·영향이 구분되게 작성한다.", size=9.5, color=GREY)
    for _ in range(3):
        para(doc, "[                                                                                    ]", space=4)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    h1(doc, "5. 조치할 사항")
    h2(doc, "① 주의")
    para(doc, "예방 중심으로 작성하고 \"앞으로 … 철저히 하시기 바랍니다\"로 맺는다.", size=9.5, color=GREY)
    para(doc, "[                                                                                    ]", space=8)
    h2(doc, "② 통보")
    para(doc, "제도 보완 중심으로 작성하고 \"…방안을 마련하시기 바랍니다\"로 맺는다.", size=9.5, color=GREY)
    para(doc, "[                                                                                    ]", space=10)

    h1(doc, "6. 작성 정보")
    blank("작성자"); blank("검토자"); blank("작성일")
    para(doc, "※ 처분요구 유형의 명칭은 공공감사에 관한 법률 제23조 제2항의 법정 유형과 기관 자체감사 규정에 맞추어 확정한다. 인용 법령의 조·항과 시행일, 수치는 공식 원문으로 대조한다.", size=9, color=GREY)

    save(doc, "KIIA-5_5개부서_감사", "audit_report_form.docx")

def kiia5():
    for fname, spec in DEPT_PLANS.items():
        kiia5_plan(fname, spec)
    kiia5_form()

# ═══════════════════════════════════════════════════════════
# KIIA-8  위험조항 삽입 용역계약서
# ═══════════════════════════════════════════════════════════
def kiia8_contract():
    doc = new_doc()
    title(doc, "소프트웨어 개발 용역계약서", "○○행정시스템 구축 용역")
    notice(doc)
    para(doc, "발주자(갑) 관점의 계약 검토 실습용 계약서다. 본문에 발주자에게 불리한 위험조항 일곱 개를 삽입하고, 각 조항 아래에 ⚠ [실습포인트] 주석을, 말미에 위험조항 목록표를 두어 정답을 표시하였다. 실전 검토에서는 정답 표시를 가린 뒤 조항을 스스로 식별한다.", size=10, color=GREY)

    para(doc, "발주자 ○○행정기관(이하 \"갑\")과 수급자 주식회사 테크솔루션즈(이하 \"을\")는 ○○행정시스템 개발 용역에 관하여 다음과 같이 계약을 체결한다.", space=8)

    def clause(no, name, body):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"제{no}조 ({name})"); _set_font(r, 11, True)
        for b in body:
            pb = doc.add_paragraph(); pb.paragraph_format.space_after = Pt(2)
            pb.paragraph_format.left_indent = Cm(0.3)
            rb = pb.add_run(b); _set_font(rb, 10.5)

    def flag(text):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.3)
        r = p.add_run("⚠ [실습포인트] "); _set_font(r, 10, True, RGBColor(0xB0,0x00,0x00))
        r2 = p.add_run(text); _set_font(r2, 10, False, RGBColor(0xB0,0x00,0x00), italic=True)

    clause("1", "목적", ["본 계약은 갑이 을에게 ○○행정시스템의 설계·개발·시험·이관 용역을 위탁하고 그 대가를 지급하는 데 필요한 사항을 정함을 목적으로 한다."])
    clause("2", "계약금액", ["계약금액은 금 일억이천만원정(₩120,000,000, 부가가치세 포함)으로 한다."])
    clause("3", "계약기간", ["계약기간은 2026년 5월 1일부터 2026년 12월 31일까지로 한다."])
    clause("4", "산출물 및 검수", ["을은 요구사항 명세서에 따른 산출물을 납품하고, 갑은 납품일로부터 14일 이내에 검수한다."])

    clause("5", "저작권", ["본 용역으로 개발된 소프트웨어 및 산출물의 저작권은 을에게 귀속한다. 갑은 을이 허락한 범위에서만 산출물을 사용할 수 있다."])
    flag("산출물 저작권이 을에게 귀속되어 발주자가 자기 비용으로 개발한 시스템을 자유롭게 사용·개작·재위탁하지 못한다. 발주자 귀속 또는 공동 귀속으로 조정이 필요하다.")

    clause("6", "대금 지급", ["갑은 검수 완료 후 30일 이내에 계약금액을 을에게 지급한다."])

    clause("7", "계약의 해지", ["을은 갑에게 30일 전 서면 통지로 본 계약을 해지할 수 있다. 갑의 해지에 관하여는 별도로 정하지 아니한다."])
    flag("해지권이 을에게만 부여되어 있고 갑의 해지 사유·절차가 규정되지 않았다. 발주자의 해지 사유(을의 채무불이행 등)와 절차를 대칭적으로 보완해야 한다.")

    clause("8", "손해배상", ["본 계약과 관련하여 을에게 발생한 손해에 대하여 갑은 그 전액을 배상한다. 을이 갑에게 배상할 책임은 계약금액의 100분의 10을 한도로 한다."])
    flag("갑의 배상책임은 한도 없이 무제한이면서 을의 배상책임은 계약금액의 10%로 제한되어 있다. 배상 한도를 대칭적으로 설정하고 무제한 배상 조항을 삭제해야 한다.")

    clause("9", "지체상금", ["갑이 대금 지급을 지연하는 경우 갑은 지연일수 1일당 계약금액의 1000분의 5에 해당하는 지연손해금을 을에게 지급한다."])
    flag("대금 지급 지연에 대한 지연손해금률(일 0.5%, 연 환산 약 182%)이 과도하다. 을의 납품 지연에 대한 지체상금 조항이 없는 점도 확인해야 한다.")

    clause("10", "계약의 갱신", ["계약기간 만료 7일 전까지 어느 당사자도 서면으로 종료를 통지하지 않으면 본 계약은 동일 조건으로 1년간 자동 갱신된다."])
    flag("자동 갱신 조항의 종료 통지 기한(만료 7일 전)이 지나치게 짧아 발주자가 갱신 여부를 검토할 시간이 부족하다. 통지 기한을 30~60일로 늘리거나 자동 갱신을 삭제해야 한다.")

    clause("11", "준거법 및 분쟁해결", ["본 계약의 준거법은 싱가포르법으로 하며, 분쟁은 싱가포르 국제중재센터의 중재로 해결한다."])
    flag("준거법과 분쟁해결지가 국내가 아닌 싱가포르로 지정되어 발주자의 분쟁 대응 비용과 부담이 크다. 준거법을 대한민국법으로, 관할을 국내 법원 또는 국내 중재로 조정해야 한다.")

    clause("12", "비밀유지", ["비밀유지 대상은 을이 서면으로 \"비밀\"이라 지정한 정보에 한정한다. 그 밖의 정보에는 비밀유지 의무가 적용되지 않는다."])
    flag("비밀유지 대상이 을이 서면 지정한 정보로 협소하게 한정되어, 발주자의 행정·개인정보 등 민감정보가 보호 범위에서 빠질 수 있다. 비밀정보의 정의를 포괄적으로 넓혀야 한다.")

    clause("13", "기타", ["본 계약에 정하지 아니한 사항은 관계 법령과 상관례에 따른다. 본 계약의 성립을 증명하기 위하여 계약서 2부를 작성해 갑과 을이 서명·날인한 후 각 1부씩 보관한다."])

    para(doc, "", space=6)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    r = p.add_run("2026년 5월 1일"); _set_font(r, 10.5)
    p = doc.add_paragraph(); r = p.add_run("갑: ○○행정기관    (인)        을: 주식회사 테크솔루션즈    (인)"); _set_font(r, 10.5)

    h1(doc, "[정답] 위험조항 목록표")
    para(doc, "아래는 본 계약서에 삽입된 발주자(갑) 관점의 위험조항 일곱 개다. 실습에서는 이 표를 가린 뒤 조항을 스스로 식별하고 대조한다.", size=9.5, color=GREY)
    table(doc, ["No", "조항", "위험 유형", "발주자 관점 문제", "검토 방향"], [
        ["1", "제5조 저작권", "저작권 귀속", "산출물 저작권이 을에게 귀속", "발주자 귀속 또는 공동 귀속"],
        ["2", "제7조 해지", "비대칭 해지권", "을만 해지 가능, 갑 해지 사유 부재", "해지 사유·절차 대칭 보완"],
        ["3", "제8조 손해배상", "무제한 손해배상", "갑 무제한·을 10% 한도", "배상 한도 대칭 설정"],
        ["4", "제10조 갱신", "자동갱신·단기통보", "종료 통지 기한 7일로 과단기", "통지 기한 연장 또는 삭제"],
        ["5", "제9조 지체상금", "과도 지연손해금", "일 0.5%(연 약 182%) 과도", "요율 인하·을 지체상금 신설"],
        ["6", "제11조 준거법", "국외 준거법·관할", "싱가포르법·싱가포르 중재", "국내법·국내 관할로 조정"],
        ["7", "제12조 비밀유지", "협소한 비밀유지", "을 지정 정보로만 한정", "비밀정보 정의 포괄 확대"],
    ], widths=[1.0, 3.0, 2.8, 5.3, 3.6], body_size=9)

    save(doc, "KIIA-8_종합실습", "sample_contract.docx")

# ═══════════════════════════════════════════════════════════
if __name__ == "__main__":
    kiia1()
    kiia3()
    kiia5()
    kiia8_contract()
    print("\nKIIA DOCX 문서 생성 완료.")
